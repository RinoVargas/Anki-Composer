import io
import requests
import openpyxl
from anki.input.input_data import InputData, InputDataRecord
from anki.input.reader.input_reader import InputReader
from compose.deck_specification import DeckSpecification


class GSheetInputReader(InputReader):

    def __init__(self, spec: DeckSpecification):
        super().__init__(spec)

    def read_input(self) -> InputData:
        data = InputData()
        records = []
        file_path = self._spec.input_config.file_path
        sheets = self._spec.input_config.sheets

        wb = _load_workbook(file_path)

        sheet_names = sheets if sheets else wb.sheetnames

        for sheet_name in sheet_names:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheet_records = list(self._parse_worksheet(ws))
                records.extend(sheet_records)

        wb.close()
        data.records = records
        return data

    def _parse_worksheet(self, ws):
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]

        for row_values in rows[1:]:
            if not any(row_values):
                continue
            row_dict = {
                headers[i]: (str(val) if val is not None else "")
                for i, val in enumerate(row_values)
                if i < len(headers)
            }
            record: InputDataRecord = self._create_record()
            for field in record.fields:
                field.field_value = row_dict.get(field.field_name, "")
            yield record


def _load_workbook(file_path: str) -> openpyxl.Workbook:
    if file_path.startswith("http://") or file_path.startswith("https://"):
        response = requests.get(file_path)
        response.raise_for_status()
        return openpyxl.load_workbook(filename=io.BytesIO(response.content), data_only=True)
    else:
        return openpyxl.load_workbook(filename=file_path, data_only=True)
