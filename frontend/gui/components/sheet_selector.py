import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from tkinter import messagebox
import threading
import io
import requests
import openpyxl
import os

class SheetSelector(ttk.LabelFrame):
    def __init__(self, master, **kwargs):
        kwargs.setdefault('text', "Select Sheets")
        kwargs.setdefault('padding', 10)
        super().__init__(master, **kwargs)
        
        self.sheet_vars = {}
        
        self.btn_load_sheets = ttk.Button(self, text="Load Sheets", bootstyle=INFO, command=self._on_load_clicked)
        self.btn_load_sheets.pack(anchor=W, pady=5)
        
        self.sheets_container = ttk.Frame(self)
        self.sheets_container.pack(fill=X, expand=True)
        
        self.current_input_type = None
        self.current_file_path = None
        
    def set_context(self, input_type: str, file_path: str):
        self.current_input_type = input_type
        self.current_file_path = file_path

    def _on_load_clicked(self):
        if not self.current_file_path:
            messagebox.showerror("Error", "Please specify an Input File Path / URL first.")
            return
        self.load_sheets(self.current_input_type, self.current_file_path)

    def load_sheets(self, input_type: str, file_path: str):
        if not file_path:
            return
            
        def fetch():
            try:
                if input_type == "gsheet" or file_path.startswith("http"):
                    response = requests.get(file_path)
                    response.raise_for_status()
                    wb = openpyxl.load_workbook(filename=io.BytesIO(response.content), data_only=True, read_only=True)
                else:
                    if not os.path.isfile(file_path):
                        raise Exception("Local file not found.")
                    wb = openpyxl.load_workbook(filename=file_path, data_only=True, read_only=True)
                    
                sheetnames = wb.sheetnames
                wb.close()
                self.after(0, lambda: self._render_sheets(sheetnames))
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Could not load sheets:\n{e}"))
                
        threading.Thread(target=fetch, daemon=True).start()
        
    def _render_sheets(self, sheetnames):
        self.clear()
        
        for i, name in enumerate(sheetnames):
            var = ttk.BooleanVar(value=True)
            cb = ttk.Checkbutton(self.sheets_container, text=name, variable=var)
            cb.grid(row=i//3, column=i%3, sticky=W, padx=10, pady=2)
            self.sheet_vars[name] = var

    def clear(self):
        for widget in self.sheets_container.winfo_children():
            widget.destroy()
        self.sheet_vars.clear()

    def get_selected_sheets(self) -> list[str]:
        return [name for name, var in self.sheet_vars.items() if var.get()]
